#!/usr/bin/env python3
"""주덕기 과장 요청 엑셀 2종 최종본 (PM 발주 2026-09-18)

첨부1  research/미청구_요청_주덕기과장_20260918.xlsx
       25년 미청구 최종 대상 8,994 중 주소 또는 변대주가 결손인 2,006건
       (주소+변대주 둘 다 결손 1,628 · 변대주만 결손 378)
       원본 = research/미청구_요청필요_20260918.json (PM 산출)

첨부2  research/불가_요청_주덕기과장_20260918.xlsx
       25년 불가에서 고압을 뺀 10,479 계기 중 주소·변대주 결손 1,376건
       원장 상태='불가' -> 고압(공종열) 제외 -> 결손 판정

★열은 다섯 개다: 계기번호 · 고객번호 · 지사 · 주소(아는 것만) · 요청항목
  MAC·시공일·앱넘버·순번은 **넣지 않는다**(영준님 2026-09-18).
  MAC 은 시공된 건에만 있는 값이라 조회키로 부적절하고, **불가 건은 MAC 칸에 계기번호가
  들어가 있다** — 불가 11,015행 중 11,009 가 MAC==계기번호이고 진짜 MAC 보유는 6건뿐이다.

★계기번호·고객번호는 셀 서식을 텍스트로 고정한다.
  엑셀이 숫자로 읽으면 앞 0 이 날아가고 영문자 접두(A0·LA)가 깨진다.

★불가 결손 판정은 **기존변대주(K열) 단독 · '0' 을 빈값** 으로 본다.
  PM 산출(1,376)과 정확히 일치하는 정의다. 변경변대주·앱변대주까지 보면 1,366 으로
  10건 줄어든다(그 10건은 교체 후 변대주를 이미 아는 건이다) — 어느 쪽을 쓸지는 PM 판단.
"""
import importlib.util
import json
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SRC1 = Path('/Users/woodelight/Projects/ami-work/research/미청구_요청필요_20260918.json')
OUT1 = ROOT / 'research/미청구_요청_주덕기과장_20260918.xlsx'
OUT2 = ROOT / 'research/불가_요청_주덕기과장_20260918.xlsx'
OUT2_JSON = ROOT / 'research/불가_요청필요_20260918.json'

COLS = ['계기번호', '고객번호', '지사', '주소', '요청항목']
WIDTHS = [16, 14, 16, 40, 14]

_s = importlib.util.spec_from_file_location(
    'bld', str(ROOT / 'scripts/build_michunggu_boost_20260918.py'))
B = importlib.util.module_from_spec(_s)
_s.loader.exec_module(B)
nm = B.norm_meter


def raw(v):
    s = str(v if v is not None else '').strip()
    return '' if s.lower() in ('', 'nan', 'none', '#n/a', 'nat') else s


def write_xlsx(path, rows, label):
    wb = Workbook()
    ws = wb.active
    ws.title = '요청'
    ws.append(COLS)
    fill = PatternFill('solid', fgColor='DDEBF7')
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(row=1, column=i)
        c.font, c.fill = Font(bold=True), fill
        c.alignment = Alignment(horizontal='center')
    for r in rows:
        ws.append([r['계기번호'], r['고객번호'], r['지사'], r['주소'], r['요청항목']])
    # ★계기번호·고객번호는 텍스트 고정 (앞 0 · 영문자 접두 보존)
    for row in ws.iter_rows(min_row=2, max_col=2):
        for c in row:
            c.number_format = '@'
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(COLS))}{ws.max_row}'
    wb.save(path)
    print(f'저장 {path}  {ws.max_row - 1:,}행', flush=True)

    # ── 원본 대조 ────────────────────────────────────────────────────────────
    chk = load_workbook(path)['요청']
    out_ids = [chk.cell(r, 1).value for r in range(2, chk.max_row + 1)]
    src_ids = [r['계기번호'] for r in rows]

    def alpha(xs):
        return sum(1 for x in xs if re.search(r'[A-Za-z]', str(x or '')))

    cust_src = sum(1 for r in rows if r['고객번호'])
    cust_out = sum(1 for r in range(2, chk.max_row + 1) if chk.cell(r, 2).value)
    addr_src = sum(1 for r in rows if r['주소'])
    addr_out = sum(1 for r in range(2, chk.max_row + 1) if chk.cell(r, 4).value)
    print(f'  [{label}] 대조 — 행수 {len(src_ids):,}/{len(out_ids):,}'
          f' · 차집합 원본에만 {len(set(src_ids)-set(out_ids))}'
          f' 엑셀에만 {len(set(out_ids)-set(src_ids))}'
          f' · 접두 {alpha(src_ids)}/{alpha(out_ids)}'
          f' · 고객번호 {cust_src:,}/{cust_out:,} · 주소 {addr_src:,}/{addr_out:,}')
    ok = (len(src_ids) == len(out_ids) and set(src_ids) == set(out_ids)
          and alpha(src_ids) == alpha(out_ids) and cust_src == cust_out
          and addr_src == addr_out)
    print(f'  판정: {"OK" if ok else "★불일치 — 보내지 마라"}')
    print(f'  요청항목: {dict(Counter(r["요청항목"] for r in rows))}')
    return ok


def build1():
    src = json.loads(SRC1.read_text())
    rows = []
    for r in src:
        has_a = bool(r['지번주소'] or r['도로명주소'])
        has_b = bool(r['변대주번호'] or r['변대주명'])
        rows.append({
            '계기번호': r['계기번호'], '고객번호': r['고객번호'], '지사': r['지사'],
            '주소': r['지번주소'] or r['도로명주소'],      # 아는 것만, 없으면 공란
            '요청항목': '주소+변대주' if not has_a and not has_b else
                     ('주소' if not has_a else '변대주'),
        })
    rows.sort(key=lambda r: (r['지사'], r['요청항목'], r['계기번호']))
    return rows


def build2():
    L = B.load_sheet(str(B.XL), 'xl/worksheets/sheet1.xml', B.LEDGER_COLS, '원장')
    bul = L[L['상태'] == '불가'].copy()
    bul['m'] = bul['계기번호'].map(nm)
    # ★고압은 **불가 행의 공종열** 기준으로 뺀다(PM 정의, 531계기 -> 10,479)
    hv = {m for m, g in zip(bul['m'], bul['공종'].astype(str)) if g == '고압'}
    recs = [r for r in bul.to_dict('records') if nm(r['계기번호']) and nm(r['계기번호']) not in hv]
    print(f'불가 행 {len(bul):,} · 고압 제외 후 행 {len(recs):,}'
          f' · 계기 {len({nm(r["계기번호"]) for r in recs}):,}')

    per = {}
    for r in recs:
        m = nm(r['계기번호'])
        d = per.setdefault(m, {'계기번호': str(r['계기번호']).strip(), '고객번호': '',
                               '지사': '', '주소': '', '변대주': ''})
        if not d['고객번호']:
            d['고객번호'] = B.norm_cust(r.get('고객번호'))
        if not d['지사']:
            d['지사'] = raw(r.get('2차사업소'))
        if not d['주소']:
            d['주소'] = raw(r.get('주소'))
        if not d['변대주']:
            v = raw(r.get('기존변대주'))       # ★K열 단독 · '0' 은 빈값
            d['변대주'] = '' if v == '0' else v

    rows = []
    for d in per.values():
        has_a, has_b = bool(d['주소']), bool(d['변대주'])
        if has_a and has_b:
            continue
        rows.append({
            '계기번호': d['계기번호'], '고객번호': d['고객번호'], '지사': d['지사'],
            '주소': d['주소'],
            '요청항목': '주소+변대주' if not has_a and not has_b else
                     ('주소' if not has_a else '변대주'),
        })
    rows.sort(key=lambda r: (r['지사'], r['요청항목'], r['계기번호']))
    OUT2_JSON.write_text(json.dumps(rows, ensure_ascii=False, indent=1))
    print(f'저장 {OUT2_JSON}  {len(rows):,}건')
    return rows


def main():
    print('=== 첨부1 미청구 ===')
    r1 = build1()
    ok1 = write_xlsx(OUT1, r1, '첨부1')
    print('\n=== 첨부2 불가 ===')
    r2 = build2()
    ok2 = write_xlsx(OUT2, r2, '첨부2')
    print(f'\n지사별(첨부2): ' + ' · '.join(
        f'{k} {v:,}' for k, v in Counter(r['지사'] for r in r2).most_common()))
    return 0 if (ok1 and ok2) else 1


if __name__ == '__main__':
    sys.exit(main())
