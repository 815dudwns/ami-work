#!/usr/bin/env python3
"""xlsx_to_db.py — 한전 엑셀을 로컬 SQLite(data/ami.db)로 적재한다.

왜 있나 (2026-09-08): 계기 하나의 시공 여부를 확인하는 데 30분이 걸린 적이 있다.
  원본 엑셀이 스크래치패드에서 지워졌고, Gmail 을 `filename:xlsx` 로만 검색해
  **zip 첨부를 놓쳤기 때문**이다(주덕기 과장 첨부는 zip 인 경우가 많다).
  영준님 지시 — "엑셀 리스트 오면 SQL 로 빨리 검색할 수 있게. 웹이 아니라 로컬 맥에."
  그래서 파일 하나(SQLite)로 끝나고, 언제든 재생성 가능한 조회용 사본을 만든다.

원본 출처
  주덕기 과장 <deokgi_7149@kdn.com> — 실효·고압·SKT·모뎀작업리스트 발주 창구.
  **첨부가 zip 인 경우가 많다.** Gmail 에서 찾을 때 `filename:xlsx` 로만 검색하지 마라.
  회수한 원본은 `data/inbox_jdg_<날짜>/` 에 그대로 둔다 — 지우지 마라.

원칙
  - 열 이름은 **원문 그대로**. 임의 개명·값 보정 금지. 정규화는 `_norm` 열로만 덧붙인다.
  - 모든 값은 TEXT 로 넣는다(계기번호 앞 0 유실 방지).
  - 같은 `(테이블, snapshot)` 재적재는 지우고 다시 넣는다(멱등). **다른 snapshot 은 누적**한다.
  - ★**한 파일 = 한 스냅샷이 깨지는 경우 snapshot 에 판 구분을 붙인다.**
    멱등키 `(테이블, snapshot)` 은 "파일 하나가 스냅샷 하나"를 전제한다. 같은 날짜의 파일이
    여러 개 들어오는 계열(예: 종로 준공·검수리스트는 한 폴더에 실효 5~14차가 같이 있다)에서는
    그 전제가 깨져서, 그대로 두면 **뒤에 적재한 파일이 앞엣것을 지운다**.
    이럴 땐 `DISCRIMINATORS` 에 파일명에서 판을 뽑는 규칙을 넣어 `20260605-실효6차` 처럼
    snapshot 을 파일 고유값으로 만들고, 그 판 이름을 `차수` 열로도 실어 정렬·필터가 되게 한다
    (영준님/PM 결정 2026-09-08).
  - 행 수가 원본과 다르면 실패다. 조용한 누락 금지.
  - 이 DB 는 **조회용 사본**이다. site-data·hapdong·jangae 빌더는 이 파일을 읽지 않는다.

사용
  python3 scripts/xlsx_to_db.py <xlsx 경로 ...>   # 지정 파일 적재
  python3 scripts/xlsx_to_db.py --all             # 아래 KNOWN_FILES 전부 재적재
  python3 scripts/xlsx_to_db.py --list            # 적재 대상 목록만 출력
"""
import argparse
import datetime as dt
import os
import re
import sqlite3
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DB_PATH = ROOT / 'data' / 'ami.db'

# ─── 적재 대상 ────────────────────────────────────────────────────────────────
# 한전·현장 데이터만. 우리가 만든 산출물(영준작업분·TBM·종로맵_6월작업분·out/)과
#   개인/사업 무관 파일은 넣지 않는다.
KNOWN_FILES = [
    # 주덕기 과장 2026-08-28 메일(zip 첨부) — awms 전 작업자 시공기록 + 장애 대상
    'data/inbox_jdg_20260828/모뎀작업리스트_20260828144931.xlsx',
    # 실효 보강현황 — 스냅샷을 누적해 리스트 간 변화를 대조한다
    'data/inbox_jdg_20260828/계기교체 보강현황_20260828.xlsx',
    'data/inbox/원본_계기교체 보강현황_20260730.xlsx',
    'data/계기교체 보강현황_20260629_1.xlsx',
    'data/boranggi-20260703.xlsx',
    'data/계기교체_보강현황_20260528.xlsx',
    'data/original/계기교체 보강현황_20260513155458.xlsx',
    # 고압철거 대상 — v1/v2/v3 를 스냅샷으로 누적
    'data/inbox_jdg_20260827/25년보강_고압철거대상_v3.xlsx',
    'data/inbox_20260813/25년보강_고압철거대상_v2.xlsx',
    'data/inbox_20260810/25년보강_고압철거대상.xlsx',
    # SKT 중계기 설치요청
    'data/inbox_jdg_20260821/SKT 중계기 AMI 모뎀 설치요청건(260814 LP 추가).xlsx',
    # DCU 간선망 해지·정지 대상 / DCU 철거 예정 개소
    'data/reference/간선망_해지_정지대상.xlsx',
    'data/reference/DCU_철거_예정_개소_목록.xlsx',
    # 종로 실효 준공·검수리스트 — 한 폴더에 실효 5~14차가 같이 있다.
    #   폴더 날짜가 전부 같아 snapshot 에 차수를 붙여 가른다(DISCRIMINATORS 참고).
    'data/종로_실효리스트_20260605/3978-2026-3032 실효5차 준공 및 검수리스트.xlsx',
    'data/종로_실효리스트_20260605/3978-2026-3033 실효6차 준공 및 검수리스트.xlsx',
    'data/종로_실효리스트_20260605/3978-2026-3034 실효7차 준공 및 검수리스트.xlsx',
    'data/종로_실효리스트_20260605/3978-2026-3035 실효8차 준공 및 검수리스트.xlsx',
    'data/종로_실효리스트_20260605/3978-2026-3150 실효12차 준공 및 검수리스트.xlsx',
    'data/종로_실효리스트_20260605/3978-2026-3151 실효13차 준공 및 검수리스트.xlsx',
    'data/종로_실효리스트_20260605/3978-2026-3153 실효14차 준공 및 검수리스트.xlsx',
]

# ─── 한 파일 = 한 스냅샷이 깨지는 계열 ───────────────────────────────────────
# (파일명 부분일치, 파일명에서 판을 뽑는 정규식). 뽑은 값을 snapshot 뒤에 붙이고
#   `차수` 열로도 싣는다. 같은 날짜 파일이 서로를 지우는 것을 막는 유일한 장치다.
DISCRIMINATORS = [
    ('준공 및 검수리스트', r'(실효\d+차)'),
]
DISCRIMINATOR_COL = '차수'

# 같은 내용이라 뺀 것(md5 동일) — 되살리려면 이유부터 확인하라
#   data/inbox_kdn/25년보강_고압철거대상_v3.xlsx        == inbox_jdg_20260827 판
#   data/inbox/간선망_해지_정지대상_20260806.xlsx        == data/reference 판
#   data/boranggi-20260629.xlsx                        계기교체 보강현황_20260629_1 과 같은 날짜(스냅샷 충돌)
# ★제주 = 사업범위 밖, 영준님 2026-09-08 삭제지시. **다시 넣지 마라.**
#   대상이던 `data/주덕기_20260605/제주지역 AMI 인프라 완전구축 Data공유_*.xlsx` 는
#   원본 폴더째 휴지통으로 옮겼다(rm 아님 — 되돌릴 수 있게 남겼다).
#   적재했던 테이블 jeju_lv_raw / jeju_print / jeju_hv_first / jeju_hv_report / jeju_hv_raw 는 DROP 했다.
#   KNOWN_FILES 에 다시 올리면 --all 이 또 적재한다 — 올리지 마라.

# ─── 시트 → 테이블 이름 ──────────────────────────────────────────────────────
# (파일명 부분일치 | None=아무 파일, 시트명 정규식, 테이블명)
SHEET_MAP = [
    (None, r'^모뎀작업리스트$', 'modem_work'),
    (None, r'^장애$', 'jangae'),
    (None, r'^계기교체[ _]보강현황', 'boranggi'),
    ('고압철거대상', r'^Sheet1$', 'gapap'),
    ('고압철거대상', r'^Sheet2$', 'gapap_sheet2'),
    ('고압철거대상', r'^주소추가$', 'gapap_addr'),
    ('SKT', r'^전체내역$', 'skt_all'),
    ('SKT', r'^작업불가$', 'skt_unable'),
    ('간선망', r'^전체DCU 현황$', 'dcu_all'),
    ('간선망', r'^검토$', 'dcu_review'),
    ('간선망', r'^요약$', 'ganseon_summary'),
    ('간선망', r'^KT모뎀', 'kt_modem_recover'),
    ('DCU_철거_예정', r'^DCU 철거 예정 개소$', 'dcu_removal'),
    ('DCU_철거_예정', r'^유지 대상 요약$', 'dcu_removal_keep'),
    ('준공 및 검수리스트', r'^준공내역서$', 'jongno_jungong'),
    ('준공 및 검수리스트', r'^검수리스트$', 'jongno_geomsu'),
]

# 담지 않는 시트 — 원본이 아니라 파생물이다.
#   모뎀작업리스트 Sheet2/Sheet3 = 엑셀 피벗 결과('행 레이블'=MAC, '개수 : 계기번호').
#   modem_work 를 GROUP BY 하면 그대로 나오고, 계기번호로 검색해도 걸리지 않는다.
SKIP_SHEETS = [
    ('모뎀작업리스트', r'^Sheet[23]$'),
]

# 인덱스를 걸 열(있을 때만)
INDEX_COLS = ['계기번호_norm', '철거계기번호_norm', 'mac_norm', '고객번호_norm', 'snapshot',
              'DCUID', 'DCU ID', '변대주번호', '지사', '작업일자',
              '차수', '차수_판']      # 판 구분(종로 실효 5~14차) 정렬·필터용

HEADER_SCAN_ROWS = 12      # 헤더를 찾을 때 훑는 상단 행 수


def log(msg):
    print(msg, flush=True)


# ─── snapshot 날짜 ───────────────────────────────────────────────────────────
def snapshot_of(path: Path) -> tuple:
    """파일명 -> YYYYMMDD. 없으면 상위 폴더명, 그래도 없으면 수정일.
    반환 (snapshot, 근거)."""
    name = path.name
    m = re.search(r'(20\d{6})', name)
    if m:
        return m.group(1), '파일명'
    m = re.search(r'\((\d{6})[^)]*\)', name)          # 예: (260814 LP 추가)
    if m:
        return '20' + m.group(1), '파일명(6자리)'
    m = re.search(r'(20\d{6})', path.parent.name)
    if m:
        return m.group(1), f'폴더명({path.parent.name})'
    ts = dt.datetime.fromtimestamp(path.stat().st_mtime)
    return ts.strftime('%Y%m%d'), '파일 수정일'


def discriminator_of(path: Path):
    """같은 날짜 파일이 여럿인 계열에서 파일을 가르는 판 이름. 해당 없으면 None."""
    for fpat, rx in DISCRIMINATORS:
        if fpat in path.name:
            m = re.search(rx, path.name)
            if m:
                return m.group(1)
            raise SystemExit(f'  실패: {path.name} — 판 구분을 뽑지 못했다(정규식 {rx}). '
                             '그대로 두면 같은 날짜 파일이 서로를 지운다')
    return None


# ─── 헤더 탐지 ───────────────────────────────────────────────────────────────
def find_header(rows, ncol):
    """헤더 행 index 를 찾는다.

    헤더의 성질: 채워진 칸이 충분히 많고(>=절반), 채워진 값이 **전부 문자열**이다.
    데이터 행에는 숫자·날짜가 섞이므로 먼저 걸리지 않는다.
    (원본_계기교체 보강현황처럼 데이터 행이 헤더보다 더 꽉 찬 경우가 있어
     '가장 꽉 찬 행'을 헤더로 잡으면 틀린다 — 그래서 '문자열만'을 함께 본다.)
    """
    need = max(2, int(ncol * 0.5))
    for i, r in enumerate(rows[:HEADER_SCAN_ROWS]):
        vals = [c for c in r if c not in (None, '')]
        if len(vals) < need:
            continue
        if all(isinstance(c, str) for c in vals):
            return i
    return None


def col_names(header_row, ncol):
    """엑셀 헤더 -> SQL 열 이름. 원문 그대로 두되 빈칸·중복만 처리."""
    out, seen = [], {}
    for i in range(ncol):
        raw = header_row[i] if i < len(header_row) else None
        name = str(raw).strip() if raw not in (None, '') else ''
        name = re.sub(r'\s+', ' ', name.replace('\n', ' ')).strip()
        if not name:
            name = f'col_{i + 1}'
        if name in seen:                    # 원본에 같은 열 이름이 두 번 나오는 파일이 있다
            seen[name] += 1
            name = f'{name}_{seen[name]}'
        else:
            seen[name] = 1
        out.append(name)
    return out


def cell_text(v):
    """모든 값을 TEXT 로. 계기번호 앞 0 을 지키려고 숫자도 문자열로 둔다."""
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        return s or None
    if isinstance(v, bool):
        return 'TRUE' if v else 'FALSE'
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return v.isoformat(sep=' ') if isinstance(v, dt.datetime) else v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


# ─── 정규화 열 ───────────────────────────────────────────────────────────────
def norm_meter(v):
    d = re.sub(r'\D', '', v or '')
    return d.zfill(11) if d else None


def norm_plain(v):
    """하이픈·공백만 제거. 앞 0 은 그대로 둔다."""
    s = re.sub(r'[\s\-]', '', v or '')
    return s.upper() or None


def pick(cols, *cands):
    """열 목록에서 첫 번째로 맞는 열 이름을 고른다(정확 일치 우선, 없으면 부분일치)."""
    for c in cands:
        if c in cols:
            return c
    for c in cands:
        for col in cols:
            if c in col:
                return col
    return None


def slug(s):
    s = re.sub(r'[^0-9A-Za-z가-힣]+', '_', str(s)).strip('_').lower()
    return s or 'sheet'


def table_for(path: Path, sheet: str, used: dict):
    """시트 -> 테이블명. 매핑표 우선, 없으면 슬러그(제네릭 이름은 파일명을 앞에 붙인다)."""
    for fpat, spat, tbl in SHEET_MAP:
        if fpat and fpat not in path.name:
            continue
        if re.search(spat, sheet):
            return tbl, '매핑표'
    # 'Sheet1' 같은 제네릭 이름은 파일마다 뜻이 달라 그대로 쓰면 섞인다
    if re.match(r'^Sheet\d+$', sheet, re.I) or slug(sheet) in used:
        return f'{slug(path.stem)[:40]}__{slug(sheet)}', '자동(파일명 결합)'
    return slug(sheet), '자동(슬러그)'


def skip_sheet(path: Path, sheet: str):
    for fpat, spat in SKIP_SHEETS:
        if fpat in path.name and re.search(spat, sheet):
            return True
    return False


# ─── 적재 ────────────────────────────────────────────────────────────────────
def ensure_table(con, table, cols):
    cur = con.execute(f'PRAGMA table_info("{table}")')
    existing = [r[1] for r in cur.fetchall()]
    if not existing:
        ddl = ', '.join(f'"{c}" TEXT' for c in cols)
        con.execute(f'CREATE TABLE "{table}" ({ddl})')
        return
    # 스냅샷마다 열 구성이 다른 파일이 있다(보강현황 24~37열) — 빠진 열만 덧붙인다
    for c in cols:
        if c not in existing:
            con.execute(f'ALTER TABLE "{table}" ADD COLUMN "{c}" TEXT')
            log(f'      + 열 추가: {c}')


def load_sheet(con, path, ws, table, snapshot, disc=None):
    rows = list(ws.iter_rows(values_only=True))
    ncol = ws.max_column
    hi = find_header(rows, ncol)
    if hi is None:
        raise SystemExit(f'  실패: [{ws.title}] 헤더 행을 찾지 못했다')
    names = col_names(rows[hi], ncol)
    unnamed = sum(1 for n in names if n.startswith('col_'))
    if unnamed > len(names) * 0.5:
        raise SystemExit(f'  실패: [{ws.title}] 헤더가 비어 있는 열이 절반을 넘는다({unnamed}/{len(names)})')
    data = rows[hi + 1:]
    data = [r for r in data if any(c not in (None, '') for c in r)]   # 완전 빈 행만 버린다
    src_rows = len(data)

    # ★'계기번호' 라는 이름이 없는 시트가 있다 — 종로 준공내역서는 신설이 '부설전력량계번호',
    #   철거가 '철거전력량계번호' 다. 이름만 보고 넘기면 계기번호로 검색해도 안 걸린다.
    meter_c = pick(names, '계기번호', '부설전력량계번호')
    remv_c = pick(names, '철거전력량계번호')
    cust_c = pick(names, '고객번호')
    mac_c = pick(names, '기존모뎀MAC', '모뎀MAC', '현재맥', 'MAC')

    extra = ['snapshot', 'src_file']
    disc_col = None
    if disc:
        # 원본에 같은 이름의 열이 이미 있으면(예: dcu_all 의 '차수') 덮지 않고 비켜 쓴다
        disc_col = DISCRIMINATOR_COL if DISCRIMINATOR_COL not in names else DISCRIMINATOR_COL + '_판'
        extra.append(disc_col)
    if meter_c:
        extra.append('계기번호_norm')
    if remv_c:
        extra.append('철거계기번호_norm')
    if mac_c:
        extra.append('mac_norm')
    if cust_c:
        extra.append('고객번호_norm')
    all_cols = names + extra

    log(f'    헤더 {hi + 1}행째 · 열 {len(names)} · 데이터 {src_rows}행 -> "{table}"')
    log(f'      정규화: 계기번호={meter_c or "-"} / 철거계기={remv_c or "-"}'
        f' / MAC={mac_c or "-"} / 고객번호={cust_c or "-"}')

    ensure_table(con, table, all_cols)
    con.execute(f'DELETE FROM "{table}" WHERE snapshot=?', (snapshot,))   # 멱등

    mi = names.index(meter_c) if meter_c else None
    ri = names.index(remv_c) if remv_c else None
    ci = names.index(cust_c) if cust_c else None
    ai = names.index(mac_c) if mac_c else None
    src = str(path.relative_to(ROOT)) if str(path).startswith(str(ROOT)) else str(path)

    payload = []
    for r in data:
        vals = [cell_text(r[i]) if i < len(r) else None for i in range(ncol)]
        rec = vals + [snapshot, src]
        if disc_col:
            rec.append(disc)
        if mi is not None:
            rec.append(norm_meter(vals[mi]))
        if ri is not None:
            rec.append(norm_meter(vals[ri]))
        if ai is not None:
            rec.append(norm_plain(vals[ai]))
        if ci is not None:
            rec.append(norm_plain(vals[ci]))
        payload.append(rec)

    ph = ','.join('?' * len(all_cols))
    cl = ','.join(f'"{c}"' for c in all_cols)
    con.executemany(f'INSERT INTO "{table}" ({cl}) VALUES ({ph})', payload)
    con.commit()

    got = con.execute(f'SELECT COUNT(*) FROM "{table}" WHERE snapshot=?', (snapshot,)).fetchone()[0]
    if got != src_rows:
        raise SystemExit(f'  실패: [{ws.title}] 행 수 불일치 — 원본 {src_rows} vs 적재 {got}')
    log(f'      적재 확인 {got}행 (원본과 일치)')
    return table, got


def add_indexes(con):
    n = 0
    for (table,) in con.execute(
            "SELECT name FROM sqlite_master WHERE type='table' "
            "AND name NOT LIKE 'sqlite_%'").fetchall():
        cols = [r[1] for r in con.execute(f'PRAGMA table_info("{table}")')]
        for c in INDEX_COLS:
            if c not in cols:
                continue
            idx = f'ix_{table}_{slug(c)}'
            con.execute(f'CREATE INDEX IF NOT EXISTS "{idx}" ON "{table}" ("{c}")')
            n += 1
    con.commit()
    return n


def load_file(con, path: Path):
    if not path.exists():
        log(f'[건너뜀] 없는 파일: {path}')
        return []
    snapshot, why = snapshot_of(path)
    disc = discriminator_of(path)
    if disc:
        snapshot = f'{snapshot}-{disc}'
        why += f' + 판 구분 {disc}(같은 날짜 파일이 여럿이라 가른다)'
    log(f'\n=== {path.name}')
    log(f'  snapshot={snapshot} ({why})')
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    used, done = {}, []
    for ws in wb.worksheets:
        if skip_sheet(path, ws.title):
            log(f'  [제외] 시트 {ws.title!r} — 피벗 집계(원본 아님)')
            continue
        if ws.max_row is None or ws.max_row < 2:
            log(f'  [제외] 시트 {ws.title!r} — 비어 있음')
            continue
        table, how = table_for(path, ws.title, used)
        used[slug(ws.title)] = table
        log(f'  시트 {ws.title!r} -> 테이블 "{table}" ({how})')
        done.append(load_sheet(con, path, ws, table, snapshot, disc))
    wb.close()
    return done


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('files', nargs='*')
    ap.add_argument('--all', action='store_true', help='KNOWN_FILES 전부 재적재')
    ap.add_argument('--list', action='store_true', help='적재 대상만 출력')
    ap.add_argument('--db', default=str(DB_PATH))
    args = ap.parse_args()

    if args.list:
        for f in KNOWN_FILES:
            p = ROOT / f
            print(('있음 ' if p.exists() else '없음 ') + f)
        return 0

    targets = [ROOT / f for f in KNOWN_FILES] if args.all else [Path(f).expanduser() for f in args.files]
    if not targets:
        ap.error('적재할 파일을 지정하거나 --all 을 써라')

    db = Path(args.db)
    db.parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(db)
    con.execute('PRAGMA journal_mode=WAL')

    total = []
    for p in targets:
        total += load_file(con, p)

    n_idx = add_indexes(con)
    con.execute('ANALYZE')
    con.commit()

    log('\n' + '=' * 60)
    log(f'DB: {db}  ({db.stat().st_size / 1e6:.1f} MB) · 인덱스 {n_idx}개')
    log('테이블별 행 수:')
    for (t,) in con.execute(
            "SELECT name FROM sqlite_master WHERE type='table' "
            "AND name NOT LIKE 'sqlite_%' ORDER BY name"):
        c = con.execute(f'SELECT COUNT(*) FROM "{t}"').fetchone()[0]
        snaps = [r[0] for r in con.execute(
            f'SELECT DISTINCT snapshot FROM "{t}" ORDER BY snapshot')]
        log(f'  {t:28} {c:>8,}행   snapshot {",".join(snaps)}')
    con.close()
    return 0


if __name__ == '__main__':
    sys.exit(main())
