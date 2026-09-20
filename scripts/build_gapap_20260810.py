#!/usr/bin/env python3
"""고압철거 데이터셋 생성 (주덕기 0810 리스트 -> data/gapap-data.json).

원본: data/inbox_20260810/25년보강_고압철거대상.xlsx (main 트리, 시트 Sheet3)
  - 행이 숨겨져 있다. read_only=False 로 열어 row_dimensions[i].hidden == False 인 행만 = 426건
필터: 주소 빈칸/#N/A 제외 -> 295, 철거구분 채워진 것(작업자철거) 제외 -> 277
좌표: 카카오 주소검색 -> 키워드검색 -> 동중심(approximate), 실패 시 네이버 폴백. 행은 절대 버리지 않는다.
스키마: data/site-data.json 원소와 동일 키(map.js 공용) + 비고(기존 스키마에 비고 키 없음)
"""
import json
import re
import sys
import threading
import collections
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import requests
from openpyxl import load_workbook

sys.path.insert(0, '/Users/woodelight/Projects/ami-work/scripts')
from _env import require_env, get_env

BASE = Path(__file__).resolve().parent.parent
XLSX = Path('/Users/woodelight/Projects/ami-work/data/inbox_20260810/25년보강_고압철거대상.xlsx')
OUT = BASE / 'data' / 'gapap-data.json'
WORKERS = 8

KAKAO_KEY = require_env('KAKAO_REST_API_KEY')
NAVER_ID = get_env('NAVER_GEOCODE_ID')
NAVER_SECRET = get_env('NAVER_GEOCODE_SECRET')

_kakao = requests.Session()
_kakao.headers.update({'Authorization': f'KakaoAK {KAKAO_KEY}'})
_naver = requests.Session()
if NAVER_ID and NAVER_SECRET:
    _naver.headers.update({
        'X-NCP-APIGW-API-KEY-ID': NAVER_ID,
        'X-NCP-APIGW-API-KEY': NAVER_SECRET,
    })

# --- 스키마 (site-data.json 키 순서 그대로) ---
BASE_KEYS = [
    '지사', '주소', '도로명주소', '계기번호', '계기타입', '고객번호', '통신방식',
    '공동주택명', '상호', '검기만료년월', '계기타입_전', '인입주', '변대주', 'DCUID',
    'lat', 'lng', '좌표정확도', 'DCU장애여부', '교체사유', '시스템등록일', '계기교체일',
    '연계수신일', '등록소요일', '사업차수_전', '통신방식_전', '검침방법_전', '검침방법',
    '사업차수', '모뎀MAC', '비고',
]

# js/utils.js parseType 과 동일. 표기는 site-data.json 라벨을 따른다.
TYPE_LABEL = {
    '17': 'E타입', '19': 'AE타입',
    '25': 'G타입', '26': 'G타입', '27': 'G타입',
    '45': 'G타입', '46': 'G타입', '47': 'G타입',
    '53': '보안계기', '55': '보안계기',
}
# 엑셀 원본 표기 -> 우리 라벨 (parseType 실패 시 폴백. js/utils.js normalizeType 과 동일 취지)
XLSX_TYPE_LABEL = {
    'E': 'E타입', 'EA': 'AE타입', 'AE': 'AE타입', 'G': 'G타입',
    'AMIGO': '보안계기', '보안계기': '보안계기',
}


def s(v):
    """셀 값 -> 트림한 문자열. None/NaN/'0' 계열은 호출부에서 따로 판단."""
    if v is None:
        return ''
    t = str(v).strip()
    return '' if t.lower() in ('none', 'nan', '#n/a') else t


def meter_no(v):
    t = s(v).replace('-', '')
    if t.endswith('.0'):
        t = t[:-2]
    return t.zfill(11)


def meter_type(no, raw):
    """계기번호 3~4자리 파싱이 1순위, 실패 시 엑셀 표기 폴백 (js/detail.js:441 과 같은 순서)."""
    t = TYPE_LABEL.get(no[2:4])
    if t:
        return t
    return XLSX_TYPE_LABEL.get(s(raw).upper(), '알수없음')


# ---------------- 좌표 ----------------

def kakao_address(query):
    if not query:
        return None
    import time as _t
    for attempt in range(4):
        try:
            r = _kakao.get('https://dapi.kakao.com/v2/local/search/address.json',
                           params={'query': query}, timeout=15)
            if r.status_code == 429:
                _t.sleep(0.5 * (attempt + 1))
                continue
            if r.status_code == 200:
                docs = r.json().get('documents', [])
                if docs:
                    d = docs[0]
                    road = (d.get('road_address') or {}).get('address_name', '')
                    jibun = (d.get('address') or {}).get('address_name', '')
                    return (road or jibun or query), float(d['y']), float(d['x'])
                return None
            _t.sleep(0.3)
        except Exception:
            _t.sleep(0.3)
    return None


def kakao_keyword(query):
    """키워드검색도 429 재시도를 넣는다 — 재시도가 없으면 병렬 실행 중 쿼터에 걸린 건이
    조용히 '못 찾음'이 되어 멀쩡한 주소가 동중심(approximate)으로 떨어진다(2026-08-10 실측)."""
    if not query:
        return None
    import time as _t
    for attempt in range(4):
        try:
            r = _kakao.get('https://dapi.kakao.com/v2/local/search/keyword.json',
                           params={'query': query, 'size': 1}, timeout=15)
            if r.status_code == 429:
                _t.sleep(0.5 * (attempt + 1))
                continue
            if r.status_code == 200:
                docs = r.json().get('documents', [])
                if docs:
                    d = docs[0]
                    return (d.get('road_address_name') or d.get('address_name') or query,
                            float(d['y']), float(d['x']))
                return None
            _t.sleep(0.3)
        except Exception:
            _t.sleep(0.3)
    return None


def naver_address(query):
    if not (query and NAVER_ID and NAVER_SECRET):
        return None
    try:
        r = _naver.get('https://naveropenapi.apigw.ntruss.com/map-geocode/v2/geocode',
                       params={'query': query}, timeout=15)
        if r.status_code == 200:
            addrs = r.json().get('addresses', [])
            if addrs:
                a = addrs[0]
                return (a.get('roadAddress') or a.get('jibunAddress') or query,
                        float(a['y']), float(a['x']))
    except Exception:
        pass
    return None


def extract_dong(address):
    m = re.match(r'(.*?[동읍면가])(\s|\d|$)', address or '')
    return m.group(1).strip() if m else None


def trim_addr(addr):
    if not addr:
        return ''
    a = re.sub(r'\s+', ' ', addr).strip()
    return re.split(r'[(]', a)[0].strip()


def resolve(addr):
    """주소 -> (좌표정확도, 도로명주소, lat, lng). 실패해도 fail 로 돌려주고 행은 살린다."""
    q = trim_addr(addr)
    for fn in (kakao_address, kakao_keyword, naver_address):
        got = fn(q)
        if got:
            return 'exact', got[0], got[1], got[2]
    # 붙어쓴 주소 보정 ("광진구 화양동111-102" -> "광진구 화양동 111-102")
    q2 = re.sub(r'([동읍면가])(\d)', r'\1 \2', q)
    if q2 != q:
        for fn in (kakao_address, kakao_keyword):
            got = fn(q2)
            if got:
                return 'exact', got[0], got[1], got[2]
    dong = extract_dong(q)
    if dong:
        got = kakao_address(dong) or kakao_keyword(dong) or naver_address(dong)
        if got:
            return 'approximate', got[0], got[1], got[2]
    return 'fail', '', None, None


# ---------------- 본체 ----------------

def load_rows():
    wb = load_workbook(XLSX)          # read_only=False — 숨김행 정보를 봐야 한다
    ws = wb['Sheet3']
    hdr = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(hdr)}
    visible = []
    for i in range(2, ws.max_row + 1):
        if ws.row_dimensions[i].hidden:
            continue
        visible.append([c.value for c in ws[i]])
    print(f'숨김 제외 표시행: {len(visible)}')

    step1 = [r for r in visible if s(r[idx['주소']])]
    print(f'주소 빈칸/#N/A 제외: {len(visible)} -> {len(step1)} (제외 {len(visible)-len(step1)})')

    step2 = [r for r in step1 if not s(r[idx['철거구분']])]
    print(f'철거구분 채워진 것 제외: {len(step1)} -> {len(step2)} (제외 {len(step1)-len(step2)})')
    return idx, step2


def main():
    idx, rows = load_rows()
    if len(rows) != 277:
        print(f'!! 기대 277건과 다름: {len(rows)}건 — 중단', file=sys.stderr)
        return 1

    recs = []
    for r in rows:
        no = meter_no(r[idx['계기번호']])
        ju2 = s(r[idx['주소2']])
        if ju2 == '0':
            ju2 = ''
        bigo = ' / '.join(x for x in (s(r[idx['비고1']]), s(r[idx['비고2']])) if x)
        rec = {k: '' for k in BASE_KEYS}
        rec.update({
            '지사': s(r[idx['2차사업소']]),
            '주소': s(r[idx['주소']]),
            '계기번호': no,
            '계기타입': meter_type(no, r[idx['계기타입']]),
            '고객번호': s(r[idx['고객번호']]),
            '통신방식': s(r[idx['통신방식']]),
            '공동주택명': ju2,
            '상호': ju2,
            '변대주': s(r[idx['앱변대주']]),
            'DCUID': s(r[idx['DCU']]),
            '교체사유': '고압철거',
            '모뎀MAC': s(r[idx['MAC']]),
            '비고': bigo,
        })
        rec['lat'] = None
        rec['lng'] = None
        recs.append(rec)

    # 주소 그룹핑 후 좌표
    grouped = collections.defaultdict(list)
    for i, rec in enumerate(recs):
        grouped[rec['주소']].append(i)
    print(f'고유 주소: {len(grouped)}건 — 좌표 조회 시작', flush=True)

    resolved = {}
    cnt = collections.Counter()
    lock = threading.Lock()
    with ThreadPoolExecutor(max_workers=WORKERS) as ex:
        futs = {ex.submit(resolve, a): a for a in grouped}
        for f in as_completed(futs):
            a = futs[f]
            resolved[a] = f.result()
            with lock:
                cnt[resolved[a][0]] += 1
                done = sum(cnt.values())
                if done % 50 == 0 or done == len(grouped):
                    print(f'  [{done}/{len(grouped)}] {dict(cnt)}', flush=True)

    for a, ids in grouped.items():
        acc, road, lat, lng = resolved[a]
        for i in ids:
            recs[i]['도로명주소'] = road
            recs[i]['lat'] = lat
            recs[i]['lng'] = lng
            recs[i]['좌표정확도'] = acc

    fails = [r for r in recs if r['좌표정확도'] == 'fail']
    if fails:
        print(f'\n!! 좌표 실패 {len(fails)}건 (행은 유지):')
        for r in fails:
            print('   ', r['계기번호'], r['지사'], r['주소'])

    OUT.write_text(json.dumps(recs, ensure_ascii=False, indent=2), encoding='utf-8')
    acc_cnt = collections.Counter(r['좌표정확도'] for r in recs)
    print(f'\n생성: {OUT} — {len(recs)}건')
    print('  좌표정확도:', dict(acc_cnt))
    print('  지사별:', dict(collections.Counter(r['지사'] for r in recs).most_common()))
    print('  계기타입:', dict(collections.Counter(r['계기타입'] for r in recs).most_common()))
    return 0


if __name__ == '__main__':
    sys.exit(main())
