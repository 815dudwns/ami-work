#!/usr/bin/env python3
"""주소 회신 요청 엑셀 생성 — 주덕기 과장 앞 (PM 지시 2026-09-18)

대상: 보정판 잔량 중 **주소 전체 결손** 4,311건
      (research/미청구_미상잔량_보정_20260918.json 의 결손필드에 '주소전체' 가 있는 것)

★열 구성은 PM 지정이다: 계기번호 · 모뎀MAC · 고객번호 · 지사 · 최종시공일 · 현재확보주소(있으면).
  고객번호는 398건뿐이지만 **한전이 가장 빨리 찾는 키**라 앞쪽에 둔다(PM 2026-09-18).
  지사별 시트로 나누지 않는다. 한 장이다.

★이 목록은 **모뎀 MAC 이 유일하게 온전한 키**다(4,290/4,311 = 99.5%).
  계기번호는 재사용되므로(오염 11~18% 실측) 계기번호만으로 조회하면 남의 개소가 나온다.
  그래서 'MAC + 계기번호 -> 설치주소' 형식으로 요청한다.

★N2(재사용) 계기 355건은 요청에서 뺀다 (PM/영준님 확인 2026-09-18).
  25년에 검기가 남은 채 철거된 계기가 26년 계기교체 때 **N2(재사용)로 다른 개소에 재투입**된
  것들이다. 그래서 그 계기번호는 지금 우리 실효 site-data / 합동 아카이브에 살아 있다.
  원장 주소에는 그 계기가 없으므로 한전에 주소를 물어봐야 **죽은 개소**를 가리킨다.
  판별 = 우리 리스트(site-data · hapdong-data-archive)에 같은 계기번호가 존재하는가.
  뺀 355건은 research/미청구_재사용제외_20260918.json 에 남긴다.

★현재확보주소 열은 **전건 공란**이다. 우리가 가진 모든 소스를 훑고도 못 찾았다는 뜻이고,
  계기번호 단독 매칭으로 나온 후보 1,587건은 오염되어 있어 **일부러 넣지 않았다**
  (영준님/PM 원칙: 빈칸이 틀린 주소보다 낫다).
"""
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / 'research/미청구_미상잔량_보정_20260918.json'
OUT = ROOT / 'research/미청구_주소요청_주덕기과장_20260918.xlsx'
OUT_REUSE = ROOT / 'research/미청구_재사용제외_20260918.json'
# N2 재사용 판별에 쓰는 우리 리스트. 합동 '지도'(hapdong-data.json)는 넣지 않는다 —
#   지금 작업 중인 개소라 재사용 판정 근거가 아니다.
REUSE_LISTS = [('실효 site-data', 'data/site-data.json'),
               ('합동 아카이브', 'data/hapdong-data-archive.json')]

# 열 순서는 PM 지정(2026-09-18 재지시): 고객번호가 한전 조회에 가장 빠르므로 앞쪽에 둔다.
COLS = ['계기번호', '모뎀MAC', '고객번호', '지사', '최종시공일', '현재확보주소(있으면)']
WIDTHS = [16, 20, 14, 16, 20, 34]


def norm_meter(v):
    """계기번호 정규화 — 영문자 접두 보존. [[meter_no_prefix_preserve]]"""
    s = re.sub(r'[\s\-]', '', str(v or '')).strip()
    if not s or s.upper() in ('NAN', 'NONE', '#N/A', '0'):
        return ''
    return s.upper() if re.search(r'[A-Za-z]', s) else s.zfill(11)


def load_reuse_index():
    idx = {}
    for name, rel in REUSE_LISTS:
        p = ROOT / rel
        if not p.exists():
            print(f'  ★{rel} 없음 — 재사용 판별이 불완전해진다')
            continue
        n = 0
        for x in json.loads(p.read_text()):
            k = norm_meter(x.get('계기번호'))
            if k and k not in idx:
                idx[k] = (name, x)
                n += 1
        print(f'  재사용 색인 {name}: {n:,}')
    return idx


def fmt_day(v):
    """원장 시공일은 '20251017163426' 과 '20260410 14:01:09' 두 꼴이 섞여 온다."""
    s = re.sub(r'\D', '', str(v or ''))
    if len(s) >= 14:
        return f'{s[0:4]}-{s[4:6]}-{s[6:8]} {s[8:10]}:{s[10:12]}:{s[12:14]}'
    if len(s) >= 8:
        return f'{s[0:4]}-{s[4:6]}-{s[6:8]}'
    return str(v or '')


def main():
    gap = json.loads(SRC.read_text())
    rows = [r for r in gap['목록'] if '주소전체' in r['결손필드']]
    print(f'주소 전체 결손 {len(rows):,}건')

    # ── N2(재사용) 계기 제외 ─────────────────────────────────────────────────
    reuse = load_reuse_index()
    excluded = []
    keep = []
    for r in rows:
        hit = reuse.get(norm_meter(r['계기번호']))
        if hit:
            name, x = hit
            excluded.append({
                '계기번호': r['계기번호'],
                '원장주소': r.get('지번주소', ''),      # 전건 공란 — 원장에 그 계기가 없다는 증거
                '우리리스트주소': x.get('지번주소') or x.get('주소') or '',
                '리스트출처': name,
                '지사': r['지사'], '구분': r['구분'], '공종': r['공종'],
                'MAC': r['MAC'], '고객번호': r['고객번호'], '최종시공일': r['최종시공일'],
            })
        else:
            keep.append(r)
    rows = keep
    print(f'N2 재사용 제외 {len(excluded):,}건 -> 요청 대상 {len(rows):,}건')
    # 지사 -> 최종시공일 순으로 정렬해 과장이 보기 쉽게 한다
    rows.sort(key=lambda r: (r['지사'], re.sub(r'\D', '', str(r['최종시공일'] or ''))))
    print(f'대상 {len(rows):,}건 · MAC 보유 {sum(1 for r in rows if r["MAC"]):,}')

    wb = Workbook()
    ws = wb.active
    ws.title = '주소요청'

    head_fill = PatternFill('solid', fgColor='DDEBF7')
    ws.append(COLS)
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True)
        c.fill = head_fill
        c.alignment = Alignment(horizontal='center')

    for r in rows:
        ws.append([
            r['계기번호'],            # ★원문 그대로 — 영문자 접두를 지우지 않는다
            r['MAC'],
            r['고객번호'],            # 있는 건 398건뿐이지만 한전은 이걸로 바로 찾는다
            r['지사'],
            fmt_day(r['최종시공일']),
            '',                       # 현재확보주소 — 전건 공란(못 찾았다는 뜻)
        ])
    # 계기번호·MAC·고객번호는 앞 0 이 날아가지 않게 텍스트로 고정
    for row in ws.iter_rows(min_row=2, max_col=3):
        for c in row:
            c.number_format = '@'

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(COLS))}{ws.max_row}'

    wb.save(OUT)
    print(f'저장 {OUT}  {ws.max_row - 1:,}행')

    from collections import Counter
    OUT_REUSE.write_text(json.dumps({
        '생성': '2026-09-18',
        '사유': ('25년에 검기가 남은 채 철거된 계기가 26년 계기교체 때 N2(재사용)로 다른 개소에 '
               '재투입됐다. 원장 주소에는 그 계기가 없으므로 한전에 주소를 물어도 죽은 개소를 '
               '가리킨다 — 주소 요청에서 빼고 우리 대상에서도 제외한다(영준님 확인 2026-09-18).'),
        '판별': '우리 리스트(data/site-data.json · data/hapdong-data-archive.json)에 같은 계기번호 존재',
        '건수': len(excluded),
        '리스트출처별': dict(Counter(e['리스트출처'] for e in excluded).most_common()),
        '구분별': dict(Counter(e['구분'] for e in excluded).most_common()),
        '지사별': dict(Counter(e['지사'] for e in excluded).most_common()),
        '목록': excluded,
    }, ensure_ascii=False, indent=1))
    print(f'저장 {OUT_REUSE}  제외 {len(excluded):,}건')

    # ── 원본 대조 ────────────────────────────────────────────────────────────
    from openpyxl import load_workbook
    chk = load_workbook(OUT)['주소요청']
    got = [(chk.cell(r, 1).value, chk.cell(r, 2).value) for r in range(2, chk.max_row + 1)]
    src_ids = [r['계기번호'] for r in rows]
    out_ids = [g[0] for g in got]
    alpha_src = sum(1 for x in src_ids if re.search(r'[A-Za-z]', str(x)))
    alpha_out = sum(1 for x in out_ids if re.search(r'[A-Za-z]', str(x)))
    cust_src = sum(1 for r in rows if r['고객번호'])
    cust_out = sum(1 for r in range(2, chk.max_row + 1) if chk.cell(r, 3).value)
    print('\n=== 원본 대조 ===')
    print(f'  행수      원본 {len(src_ids):,} · 엑셀 {len(out_ids):,}')
    print(f'  고객번호  원본 {cust_src:,} · 엑셀 {cust_out:,}')
    print(f'  차집합    원본에만 {len(set(src_ids)-set(out_ids))} · 엑셀에만 {len(set(out_ids)-set(src_ids))}')
    print(f'  접두 문자 원본 {alpha_src} · 엑셀 {alpha_out}')
    print(f'  MAC 공란  {sum(1 for g in got if not g[1]):,}')
    ok = (len(src_ids) == len(out_ids) and set(src_ids) == set(out_ids)
          and alpha_src == alpha_out and cust_src == cust_out)
    print('  판정:', 'OK' if ok else '★불일치 — 배포하지 마라')
    return 0 if ok else 1


if __name__ == '__main__':
    sys.exit(main())
